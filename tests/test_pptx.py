import os
import sys

# Ensure root in sys.path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from src.data_loader import DataLoader
from src.ppt_generator import PowerPointGenerator


def verify_pptx(pptx_path):
    print(f"Verifying {pptx_path}...")
    prs = Presentation(pptx_path)
    print(f"Slide width: {prs.slide_width.inches:.3f} in, Slide height: {prs.slide_height.inches:.3f} in")
    print(f"Total slides: {len(prs.slides)}")
    assert len(prs.slides) >= 2, f"Expected at least 2 slides, got {len(prs.slides)}"

    # Check title/summary slide
    title_slide = prs.slides[0]
    has_title = any(
        "KPI Benchmark" in s.text_frame.text
        for s in title_slide.shapes
        if s.has_text_frame
    )
    assert has_title, "Title slide does not have 'KPI Benchmark' title"

    # Check KPI slides
    for idx in range(1, len(prs.slides)):
        slide = prs.slides[idx]
        text_shapes = [s for s in slide.shapes if s.has_text_frame]
        title_box = text_shapes[0]
        title_para = title_box.text_frame.paragraphs[0]
        align = title_para.alignment
        print(f"Slide {idx+1} title: '{title_para.text}', Alignment: {align} (PP_ALIGN.CENTER={PP_ALIGN.CENTER})")
        assert align == PP_ALIGN.CENTER, f"Slide {idx+1} title is not center-aligned!"

    print("ALL PPTX VERIFICATIONS PASSED!\n")


if __name__ == "__main__":
    if len(sys.argv) > 1 and os.path.exists(sys.argv[1]):
        target = sys.argv[1]
    else:
        target = "output/test_cli_verification.pptx"
        sample_csv = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "sample_network_kpis.csv")
        loader = DataLoader(sample_csv)
        sug = loader.get_suggested_periods()
        gen = PowerPointGenerator()
        gen.generate_report(
            loader=loader,
            selected_kpis=loader.numeric_kpis[:3],
            pre_start=sug["pre_start"],
            pre_end=sug["pre_end"],
            impl_date=sug["impl_date"],
            post_start=sug["post_start"],
            post_end=sug["post_end"],
            output_path=target,
            include_summary_table=True,
        )

    verify_pptx(target)
