import os
import sys

# Ensure root in sys.path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from src.data_loader import DataLoader
from src.ppt_generator import PowerPointGenerator


def test_summary_table_generation():
    print("=== Testing Summary Table on Slide 1 ===")
    sample_file = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "sample_network_kpis.csv")
    loader = DataLoader(sample_file)
    sug = loader.get_suggested_periods()

    selected_kpis = loader.numeric_kpis[:8]

    out_path = "output/test_summary_table_verified.pptx"
    gen = PowerPointGenerator()
    gen.generate_report(
        loader=loader,
        selected_kpis=selected_kpis,
        pre_start=sug["pre_start"],
        pre_end=sug["pre_end"],
        impl_date=sug["impl_date"],
        post_start=sug["post_start"],
        post_end=sug["post_end"],
        output_path=out_path,
        include_summary_table=True,
    )

    prs = Presentation(out_path)
    total_slides = len(prs.slides)
    expected_slides = 1 + len(selected_kpis)
    print(f"Total slides generated: {total_slides} (Expected: 1 summary slide + {len(selected_kpis)} KPI slides = {expected_slides})")
    assert total_slides == expected_slides, f"Expected {expected_slides} slides, got {total_slides}"

    # Slide 1: Summary Table
    slide_1 = prs.slides[0]
    tables = [s.table for s in slide_1.shapes if s.has_table]
    assert len(tables) == 1, "Slide 1 must have exactly 1 native Table shape"
    table = tables[0]

    # Verify table dimensions
    assert len(table.columns) == 4, f"Expected 4 columns, got {len(table.columns)}"
    assert len(table.rows) == len(selected_kpis) + 1, f"Expected {len(selected_kpis)+1} rows, got {len(table.rows)}"

    # Check Headers
    headers = [table.cell(0, i).text_frame.text.strip() for i in range(4)]
    print(f"Slide 1 Table Headers: {headers}")
    assert headers == ["KPIs", "Pre", "Post", "Change%"], f"Unexpected headers: {headers}"

    # Check Data rows formatting and colors
    for r_idx, kpi in enumerate(selected_kpis, 1):
        cell_kpi = table.cell(r_idx, 0)
        cell_pre = table.cell(r_idx, 1)
        cell_post = table.cell(r_idx, 2)
        cell_chg = table.cell(r_idx, 3)

        assert cell_kpi.text_frame.text == kpi, f"Row {r_idx} KPI name mismatch"
        assert cell_pre.text_frame.text != "", f"Row {r_idx} Pre value empty"
        assert cell_post.text_frame.text != "", f"Row {r_idx} Post value empty"
        assert cell_chg.text_frame.text != "", f"Row {r_idx} Change% value empty"

        # Check column 0 fill is blue, column 1 is white
        assert cell_kpi.fill.fore_color.rgb == RGBColor(207, 226, 243)
        assert cell_pre.fill.fore_color.rgb == RGBColor(255, 255, 255)

    print("[PASS] Slide 1 Table structure, styling, and data verified!")

    # Verify KPI slides follow
    for s_idx in range(1, len(prs.slides)):
        kpi_slide = prs.slides[s_idx]
        title_box = [s for s in kpi_slide.shapes if s.has_text_frame][0]
        title_text = title_box.text_frame.paragraphs[0].text
        align = title_box.text_frame.paragraphs[0].alignment
        assert align == PP_ALIGN.CENTER, f"Slide {s_idx+1} title '{title_text}' is not center-aligned"

    print("[PASS] KPI slides center-alignment verified!")


def test_summary_table_pagination():
    print("\n=== Testing Pagination with 20 KPIs ===")
    sample_file = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "sample_network_kpis.csv")
    loader = DataLoader(sample_file)
    sug = loader.get_suggested_periods()

    # Repeat KPIs to create 20 entries
    selected_20 = (loader.numeric_kpis * 2)[:20]
    out_path = "output/test_summary_table_paginated.pptx"
    gen = PowerPointGenerator()
    gen.generate_report(
        loader=loader,
        selected_kpis=selected_20,
        pre_start=sug["pre_start"],
        pre_end=sug["pre_end"],
        impl_date=sug["impl_date"],
        post_start=sug["post_start"],
        post_end=sug["post_end"],
        output_path=out_path,
        include_summary_table=True,
    )

    prs = Presentation(out_path)
    # 20 KPIs split into chunk of 16 + chunk of 4 = 2 summary slides + 20 KPI slides = 22 slides
    print(f"Total slides for 20 KPIs: {len(prs.slides)} (Expected: 2 summary slides + 20 KPI slides = 22)")
    assert len(prs.slides) == 22, f"Expected 22 slides, got {len(prs.slides)}"

    # Check Slide 1 table rows
    s1_table = [s.table for s in prs.slides[0].shapes if s.has_table][0]
    assert len(s1_table.rows) == 17, f"Slide 1 table should have 1 header + 16 data rows = 17, got {len(s1_table.rows)}"

    # Check Slide 2 table rows
    s2_table = [s.table for s in prs.slides[1].shapes if s.has_table][0]
    assert len(s2_table.rows) == 5, f"Slide 2 table should have 1 header + 4 data rows = 5, got {len(s2_table.rows)}"

    print("[PASS] Pagination verified cleanly!")


if __name__ == "__main__":
    test_summary_table_generation()
    test_summary_table_pagination()
    print("\nALL SUMMARY TABLE TESTS PASSED!")
